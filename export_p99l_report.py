import os
import sys
libs_path = os.path.join(os.path.dirname(__file__), 'libs')
if os.path.exists(libs_path) and libs_path not in sys.path:
    sys.path.insert(0, libs_path)

import json
import urllib.request
import urllib.parse
import ssl
from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ssl_ctx = ssl._create_unverified_context()

def format_number(val):
    if val is None or val == 0:
        return 0
    if isinstance(val, float) and val.is_integer():
        return int(val)
    return val

def parse_input_date(date_str):
    if not date_str:
        return datetime.now().strftime("%Y-%m-%d")
    date_str = date_str.strip().lower()
    if date_str == "today":
        return datetime.now().strftime("%Y-%m-%d")
    if date_str == "yesterday":
        return (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
    
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d_%m_%Y"]:
        try:
            return datetime.strptime(date_str, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"Unknown date format: {date_str}. Please use YYYY-MM-DD or DD/MM/YYYY.")

LOTTERY_CONFIGS = [
    {"sheet": "Data_MHSB", "code": "MHSB", "api_code": "LEAP", "name": "លាភមហាសម្បត្តិ (MHSB)"},
    {"sheet": "Data_MC",   "code": "MC",   "api_code": "VN1",  "name": "មហាឈ្នះ (MC)"},
    {"sheet": "Data_MT",   "code": "MT",   "api_code": "VN2",  "name": "មហាទេព (MT)"},
    {"sheet": "Data_KH",   "code": "KH",   "api_code": "KH",   "name": "ឆ្នោតខ្មែរ (KH)"},
    {"sheet": "Data_TC",   "code": "TC",   "api_code": "TN",   "name": "ទិញឈ្នះ (TC)"},
    {"sheet": "Data_SC",   "code": "SC",   "api_code": "SC",   "name": "ឆ្នោត (SC)"},
    {"sheet": "Data_TH",   "code": "TH",   "api_code": "TH",   "name": "ឆ្នោត (TH)"},
    {"sheet": "Data_KP",   "code": "KP",   "api_code": "KP",   "name": "ឆ្នោត កំពត (KP)"}
]

HEADERS_32 = [
    "Date", "ប្រភេទឆ្នោត", "កូដសមាជិក", "ឈ្មោះ",
    "លុយលក់ 1D(KHR)", "លុយលក់ 2D(KHR)", "លុយលក់ 3D(KHR)", "លុយលក់ 4D(KHR)",
    "លុយលក់ 1D(USD)", "លុយលក់ 2D(USD)", "លុយលក់ 3D(USD)", "លុយលក់ 4D(USD)",
    "លុយសង 1D(KHR)", "លុយសង 2D(KHR)", "លុយសង 3D(KHR)", "លុយសង 4D(KHR)",
    "លុយសង 1D(USD)", "លុយសង 2D(USD)", "លុយសង 3D(USD)", "លុយសង 4D(USD)",
    "លុយស៊ីខាត (KHR)", "លុយស៊ីខាត (USD)",
    "បញ្ជីចាស់ (KHR)", "បញ្ជីចាស់ (USD)",
    "លុយខ្លី (KHR)", "លុយសង (KHR)",
    "លុយខ្លី (USD)", "លុយសង (USD)",
    "តវ៉ា (KHR)", "តវ៉ា (USD)",
    "ប្រាក់តុល្យភាពចុងក្រោយ (KHR)", "ប្រាក់តុល្យភាពចុងក្រោយ (USD)"
]

def fetch_report_data(base_url, endpoint_path, auth_headers, level, date_str):
    params = {
        "filterByLevel": level,
        "filterByReportType": "TOTAL",
        "startDate": date_str,
        "endDate": date_str,
        "filterByUsername": "",
        "size": 10000
    }
    query_str = urllib.parse.urlencode(params)
    url = f"{base_url}{endpoint_path}?{query_str}"
    
    req = urllib.request.Request(url, headers=auth_headers, method="GET")
    try:
        with urllib.request.urlopen(req, context=ssl_ctx, timeout=15) as resp:
            res_json = json.loads(resp.read().decode('utf-8'))
            return res_json.get('data', {})
    except Exception as e:
        print(f"❌ Failed to fetch {level}: {e}")
        return {}

def fetch_single_lottery_data(base_url, auth_headers, api_code, date_str):
    url = f"{base_url}/api/v1.0.0/total-report?filterByLevel=member&filterByReportType=TOTAL&filterByLotteryType={api_code}&startDate={date_str}&endDate={date_str}&size=10000"
    req = urllib.request.Request(url, headers=auth_headers, method="GET")
    try:
        with urllib.request.urlopen(req, context=ssl_ctx, timeout=15) as resp:
            res_json = json.loads(resp.read().decode('utf-8'))
            return res_json.get('data', {})
    except Exception as e:
        return {}

def append_rows_to_sheet(ws, date_display_str, lottery_label, settlements, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right):
    start_row = ws.max_row + 1
    for r_offset, item in enumerate(settlements):
        curr_row = start_row + r_offset
        ws.row_dimensions[curr_row].height = 19
        u_name = item.get('username') or item.get('userCode') or ''
        nick = item.get('nickname') or ''

        row_vals = [
            date_display_str,
            lottery_label,
            u_name,
            nick,
            item.get('betAmount1DKhr', 0),
            item.get('betAmount2DKhr', 0),
            item.get('betAmount3DKhr', 0),
            item.get('betAmount4DKhr', 0),
            item.get('betAmount1DUsd', 0),
            item.get('betAmount2DUsd', 0),
            item.get('betAmount3DUsd', 0),
            item.get('betAmount4DUsd', 0),
            item.get('winAmount1DKhr', 0),
            item.get('winAmount2DKhr', 0),
            item.get('winAmount3DKhr', 0),
            item.get('winAmount4DKhr', 0),
            item.get('winAmount1DUsd', 0),
            item.get('winAmount2DUsd', 0),
            item.get('winAmount3DUsd', 0),
            item.get('winAmount4DUsd', 0),
            item.get('winLoseAmountKhr', 0),
            item.get('winLoseAmountUsd', 0),
            item.get('oldAmountKhr', 0),
            item.get('oldAmountUsd', 0),
            (item.get('borrow') or {}).get('amountKhr', 0),
            (item.get('give') or {}).get('amountKhr', 0),
            (item.get('borrow') or {}).get('amountUsd', 0),
            (item.get('give') or {}).get('amountUsd', 0),
            (item.get('protestAmount') or {}).get('amountKhr', 0),
            (item.get('protestAmount') or {}).get('amountUsd', 0),
            item.get('totalAmountKhr', 0),
            item.get('totalAmountUsd', 0)
        ]

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = cell_font
            cell.border = grid_border
            if c_idx == 1:
                cell.alignment = align_center
            elif c_idx == 2:
                cell.alignment = align_center
                cell.font = blue_font
            elif c_idx in [3, 4]:
                cell.alignment = align_left
                cell.font = blue_font
            else:
                cell.alignment = align_right
                if isinstance(val, (int, float)):
                    cell.number_format = '#,##0'
                    if val < 0:
                        cell.font = red_bold_font

def update_daily_report(raw_date=None):
    target_date_iso = parse_input_date(raw_date)
    dt_obj = datetime.strptime(target_date_iso, "%Y-%m-%d")
    target_date_display = dt_obj.strftime("%d/%m/%Y")

    print(f"\n=======================================================")
    print(f"🚀 Updating Master Database for Date: {target_date_iso} ({target_date_display})")
    print(f"=======================================================\n")

    # Load credentials
    env = {}
    env_file = os.path.join(os.getcwd(), '.env')
    if os.path.exists(env_file):
        with open(env_file, 'r', encoding='utf-8') as f:
            for line in f:
                if '=' in line:
                    k, v = line.strip().split('=', 1)
                    env[k.strip()] = v.strip().strip('\"').strip('\'')

    username = env.get('P99L_USERNAME')
    password = env.get('P99L_PASSWORD')

    base_url = "https://d1rsx2ylunqjey.cloudfront.net"
    login_headers = {
        'Content-Type': 'application/json',
        'Authorization': 'Basic R2VuaXVzTG9UVEVSWTE2OEAxMjNAIzIlOTk5OV4mJDpMb3R0b1YyMTY4RnJvbnRAMTIzQCMyJTk5OTleJiQ=',
        'Origin': 'https://admin.p99l.com'
    }

    login_req = urllib.request.Request(f"{base_url}/api/v1.0.0/auth/login", data=json.dumps({'username': username, 'password': password}).encode('utf-8'), headers=login_headers, method='POST')
    with urllib.request.urlopen(login_req, context=ssl_ctx, timeout=15) as resp:
        res = json.loads(resp.read().decode('utf-8'))
        token = res.get('data', {}).get('token')
        user_info = res.get('data', {}).get('user', {})

    auth_headers = {
        'Authorization': f'Bearer {token}',
        'Origin': 'https://admin.p99l.com',
        'Referer': 'https://admin.p99l.com/'
    }
    print(f"🔑 Logged in successfully as {user_info.get('username')}")

    # 1. Fetch 4 user levels for ALL lotteries
    print("📡 Fetching Member, Master, Senior, SuperSenior (ALL)...")
    member_data = fetch_report_data(base_url, "/api/v1.0.0/report/total-report/all", auth_headers, "member", target_date_iso)
    master_data = fetch_report_data(base_url, "/api/v1.0.0/report/total-report/all", auth_headers, "master", target_date_iso)
    senior_data = fetch_report_data(base_url, "/api/v1.0.0/report/total-report/all", auth_headers, "senior", target_date_iso)
    super_senior_data = fetch_report_data(base_url, "/api/v1.0.0/report/total-report/all", auth_headers, "super-senior", target_date_iso)

    # 2. Fetch 8 individual lotteries for Member & SuperSenior
    lottery_sales_summary = {}
    lottery_settlements_map = {}
    super_senior_lotteries = {}
    print("\n📡 Fetching 8 Individual Lottery Types...")
    for cfg in LOTTERY_CONFIGS:
        l_code = cfg["code"]
        api_code = cfg["api_code"]
        data_obj = fetch_single_lottery_data(base_url, auth_headers, api_code, target_date_iso)
        summ = data_obj.get('summery', {})
        sett = data_obj.get('settlements', [])
        
        tot_khr = (summ.get('betAmount1DKhr',0) or 0) + (summ.get('betAmount2DKhr',0) or 0) + (summ.get('betAmount3DKhr',0) or 0) + (summ.get('betAmount4DKhr',0) or 0)
        tot_usd = (summ.get('betAmount1DUsd',0) or 0) + (summ.get('betAmount2DUsd',0) or 0) + (summ.get('betAmount3DUsd',0) or 0) + (summ.get('betAmount4DUsd',0) or 0)
        net_khr = summ.get('winLoseAmountKhr', 0) or 0
        net_usd = summ.get('winLoseAmountUsd', 0) or 0

        lottery_sales_summary[l_code] = {
            "sales_khr": tot_khr,
            "sales_usd": tot_usd,
            "net_khr": net_khr,
            "net_usd": net_usd,
            "b1k": summ.get('betAmount1DKhr',0) or 0,
            "b2k": summ.get('betAmount2DKhr',0) or 0,
            "b3k": summ.get('betAmount3DKhr',0) or 0,
            "b4k": summ.get('betAmount4DKhr',0) or 0,
            "b1u": summ.get('betAmount1DUsd',0) or 0,
            "b2u": summ.get('betAmount2DUsd',0) or 0,
            "b3u": summ.get('betAmount3DUsd',0) or 0,
            "b4u": summ.get('betAmount4DUsd',0) or 0
        }
        lottery_settlements_map[l_code] = sett

        # Fetch SuperSenior level for this lottery
        try:
            ss_url = f"{base_url}/api/v1.0.0/total-report?filterByLevel=super-senior&filterByReportType=TOTAL&filterByLotteryType={api_code}&startDate={target_date_iso}&endDate={target_date_iso}&size=1000"
            ss_req = urllib.request.Request(ss_url, headers=auth_headers, method="GET")
            with urllib.request.urlopen(ss_req, context=ssl_ctx, timeout=10) as resp:
                ss_res = json.loads(resp.read().decode('utf-8'))
                ss_sett = ss_res.get('data', {}).get('settlements', [])
                ss_map = {}
                for item in ss_sett:
                    u = item.get('username') or item.get('userCode')
                    if u:
                        ss_map[u] = {
                            'username': u,
                            'nickname': item.get('nickname', ''),
                            'b1k': item.get('betAmount1DKhr', 0) or 0,
                            'b2k': item.get('betAmount2DKhr', 0) or 0,
                            'b3k': item.get('betAmount3DKhr', 0) or 0,
                            'b4k': item.get('betAmount4DKhr', 0) or 0,
                            'b1u': item.get('betAmount1DUsd', 0) or 0,
                            'b2u': item.get('betAmount2DUsd', 0) or 0,
                            'b3u': item.get('betAmount3DUsd', 0) or 0,
                            'b4u': item.get('betAmount4DUsd', 0) or 0,
                            'sales_khr': (item.get('betAmount1DKhr', 0) or 0) + (item.get('betAmount2DKhr', 0) or 0) + (item.get('betAmount3DKhr', 0) or 0) + (item.get('betAmount4DKhr', 0) or 0),
                            'sales_usd': (item.get('betAmount1DUsd', 0) or 0) + (item.get('betAmount2DUsd', 0) or 0) + (item.get('betAmount3DUsd', 0) or 0) + (item.get('betAmount4DUsd', 0) or 0),
                            'net_khr': item.get('winLoseAmountKhr', 0) or 0,
                            'net_usd': item.get('winLoseAmountUsd', 0) or 0,
                            'total_khr': item.get('totalAmountKhr', 0) or 0,
                            'total_usd': item.get('totalAmountUsd', 0) or 0
                        }
                super_senior_lotteries[l_code] = ss_map
        except Exception as e:
            print(f"  ⚠️ Failed fetching SuperSenior {l_code}: {e}")

        print(f"  ✅ {cfg['name']}: Sales={tot_khr:,} ៛ | Net={net_khr:,} ៛")

    # 2.5 Fetch SuperSenior Sang Data (របាយការណ៍សាង -> របាយការណ៍សរុបរួម សាង)
    print("\n📡 Fetching SuperSenior Sang Report (លេខសាង)...")
    super_senior_sang_data = fetch_report_data(base_url, "/api/v1.0.0/send/total-report/all", auth_headers, "super-senior", target_date_iso)
    sang_summ = super_senior_sang_data.get('summery', {})
    sang_b_khr = (sang_summ.get('betAmount1DKhr',0) or 0) + (sang_summ.get('betAmount2DKhr',0) or 0) + (sang_summ.get('betAmount3DKhr',0) or 0) + (sang_summ.get('betAmount4DKhr',0) or 0)
    sang_net_khr = sang_summ.get('winLoseAmountKhr', 0)
    print(f"  ✅ លេខសាង (SuperSenior): Sales={sang_b_khr:,} ៛ | Net={sang_net_khr:,} ៛")

    # 3. Update total_summary_report.json
    main_json_path = os.path.join(os.getcwd(), 'total_summary_report.json')
    if os.path.exists(main_json_path):
        with open(main_json_path, 'r', encoding='utf-8') as f:
            main_json = json.load(f)
    else:
        main_json = {"reports": {}, "dates_available": []}

    reports_obj = main_json.get('reports', {})
    reports_obj[target_date_display] = {
        "date": target_date_display,
        "member": member_data,
        "master": master_data,
        "senior": senior_data,
        "super_senior": super_senior_data,
        "super_senior_sang": super_senior_sang_data,
        "lottery_sales_summary": lottery_sales_summary,
        "super_senior_lotteries": super_senior_lotteries
    }

    # Month detection & archiving
    target_dt = datetime.strptime(target_date_iso, "%Y-%m-%d")
    target_month_num = target_dt.strftime("%m")
    target_year_str = target_dt.strftime("%Y")

    MONTH_MAP = {
        '01': ('january', 'មករា', 'JANUARY'),
        '02': ('february', 'កុម្ភៈ', 'FEBRUARY'),
        '03': ('march', 'មីនា', 'MARCH'),
        '04': ('april', 'មេសា', 'APRIL'),
        '05': ('may', 'ឧសភា', 'MAY'),
        '06': ('june', 'មិថុនា', 'JUNE'),
        '07': ('july', 'កក្កដា', 'JULY'),
        '08': ('august', 'សីហា', 'AUGUST'),
        '09': ('september', 'កញ្ញា', 'SEPTEMBER'),
        '10': ('october', 'តុលា', 'OCTOBER'),
        '11': ('november', 'វិច្ឆិកា', 'NOVEMBER'),
        '12': ('december', 'ធ្នូ', 'DECEMBER')
    }
    m_en, m_kh, m_upper = MONTH_MAP.get(target_month_num, ('september', 'កញ្ញា', 'SEPTEMBER'))

    # If active JSON contains only previous month dates, archive it and start fresh for new month
    existing_dates = main_json.get('dates_available', [])
    if existing_dates:
        prev_month_num = existing_dates[0].split('/')[1] if len(existing_dates[0].split('/')) == 3 else None
        if prev_month_num and prev_month_num != target_month_num:
            os.makedirs('archive', exist_ok=True)
            prev_m_en = MONTH_MAP.get(prev_month_num, ('month', '', ''))[0]
            archive_json_path = f'archive/total_summary_report_{prev_m_en}_{target_year_str}.json'
            with open(archive_json_path, 'w', encoding='utf-8') as af:
                json.dump(main_json, af, ensure_ascii=False, indent=2)
            print(f"📦 Archived previous month ({prev_m_en}) dataset to {archive_json_path}")
            # Filter reports_obj for current month
            reports_obj = {k: v for k, v in reports_obj.items() if k.endswith(f"/{target_month_num}/{target_year_str}")}

    sorted_dates = sorted(list(reports_obj.keys()), key=lambda x: datetime.strptime(x, "%d/%m/%Y"), reverse=True)
    main_json["reports"] = reports_obj
    main_json["dates_available"] = sorted_dates
    main_json["date"] = sorted_dates[0] if sorted_dates else target_date_display
    main_json["month"] = f"{target_month_num}-{target_year_str}"
    main_json["month_name"] = f"{m_kh} {target_year_str} ({m_upper} {target_year_str})"
    main_json["member"] = reports_obj[sorted_dates[0]].get('member', {}) if sorted_dates else {}
    main_json["master"] = reports_obj[sorted_dates[0]].get('master', {}) if sorted_dates else {}
    main_json["senior"] = reports_obj[sorted_dates[0]].get('senior', {}) if sorted_dates else {}
    main_json["super_senior"] = reports_obj[sorted_dates[0]].get('super_senior', {}) if sorted_dates else {}

    with open(main_json_path, 'w', encoding='utf-8') as f:
        json.dump(main_json, f, ensure_ascii=False, indent=2)
    print(f"\n✅ Updated total_summary_report.json with date {target_date_display} for month {m_upper} {target_year_str}")

    # 4. Rebuild / Update Master Excel Database for Active Month
    master_excel_filename = f"p99l_master_database_{m_en}_{target_year_str}.xlsx"
    master_excel_path = os.path.join(os.getcwd(), master_excel_filename)
    
    # We run full rebuild to maintain chronological sorting & clean formatting
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    green_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    dark_green_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    header_font = Font(name="Kantumruy Pro", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Kantumruy Pro", size=14, bold=True, color="FFFFFF")
    cell_font = Font(name="Kantumruy Pro", size=9)
    blue_font = Font(name="Kantumruy Pro", size=9, bold=True, color="0046FF")
    red_bold_font = Font(name="Kantumruy Pro", size=9, bold=True, color="D93025")
    thin_border = Side(style='thin', color='B0BEC5')
    grid_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    sorted_dates_chrono = sorted(list(reports_obj.keys()), key=lambda x: datetime.strptime(x, "%d/%m/%Y"))

    # Sheet: Dashboard
    ws_dash = wb.create_sheet(title="Dashboard")
    ws_dash.views.sheetView[0].showGridLines = True
    ws_dash.merge_cells("A1:K2")
    ws_dash["A1"] = f"P99L LOTTERY SYSTEM — MASTER DASHBOARD & DAILY SUMMARY ({m_upper} {target_year_str})"
    ws_dash["A1"].fill = dark_green_fill
    ws_dash["A1"].font = title_font
    ws_dash["A1"].alignment = align_center

    month_sales_khr, month_net_khr = 0, 0
    for d_str in sorted_dates_chrono:
        rep = reports_obj.get(d_str, {})
        summ = rep.get('member', {}).get('summery', {})
        b_k = (summ.get('betAmount1DKhr',0) or 0) + (summ.get('betAmount2DKhr',0) or 0) + (summ.get('betAmount3DKhr',0) or 0) + (summ.get('betAmount4DKhr',0) or 0)
        n_k = summ.get('winLoseAmountKhr',0) or 0
        month_sales_khr += b_k
        month_net_khr += n_k

    ws_dash.merge_cells("A4:C4")
    ws_dash["A4"] = "💰 លុយលក់សរុបប្រចាំខែ (Total Sales)"
    ws_dash["A4"].fill = green_fill
    ws_dash["A4"].font = header_font
    ws_dash["A4"].alignment = align_center
    ws_dash.merge_cells("A5:C5")
    ws_dash["A5"] = f"{round(month_sales_khr):,} ៛"
    ws_dash["A5"].font = Font(name="Kantumruy Pro", size=14, bold=True, color="2E7D32")
    ws_dash["A5"].alignment = align_center

    ws_dash.merge_cells("E4:G4")
    ws_dash["E4"] = "⚖️ ស៊ីខាតសរុបប្រចាំខែ (Total Net Win/Loss)"
    ws_dash["E4"].fill = green_fill
    ws_dash["E4"].font = header_font
    ws_dash["E4"].alignment = align_center
    ws_dash.merge_cells("E5:G5")
    ws_dash["E5"] = f"{round(month_net_khr):,} ៛"
    ws_dash["E5"].font = Font(name="Kantumruy Pro", size=14, bold=True, color="2E7D32" if month_net_khr >= 0 else "D93025")
    ws_dash["E5"].alignment = align_center

    ws_dash.merge_cells("I4:K4")
    ws_dash["I4"] = "📅 ចំនួនថ្ងៃសរុប (Total Days)"
    ws_dash["I4"].fill = green_fill
    ws_dash["I4"].font = header_font
    ws_dash["I4"].alignment = align_center
    ws_dash.merge_cells("I5:K5")
    ws_dash["I5"] = f"{len(sorted_dates_chrono)} ថ្ងៃ"
    ws_dash["I5"].font = Font(name="Kantumruy Pro", size=14, bold=True, color="0046FF")
    ws_dash["I5"].alignment = align_center

    ws_dash.merge_cells("A7:K7")
    ws_dash["A7"] = "📊 តារាងសរុបស៊ីខាតប្រចាំថ្ងៃ & លុយលក់តាមប្រភេទឆ្នោតនីមួយៗ"
    ws_dash["A7"].fill = dark_green_fill
    ws_dash["A7"].font = header_font
    ws_dash["A7"].alignment = align_left

    dash_headers = ["កាលបរិច្ឆេទ", "MHSB (៛)", "MC (៛)", "MT (៛)", "KH (៛)", "TC (៛)", "SC (៛)", "TH (៛)", "KP (៛)", "លុយលក់សរុប (៛)", "ស៊ីខាតប្រចាំថ្ងៃ (៛)"]
    ws_dash.row_dimensions[8].height = 24
    for c_i, h_t in enumerate(dash_headers, 1):
        cell = ws_dash.cell(row=8, column=c_i, value=h_t)
        cell.fill = green_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = grid_border

    for r_idx, d_str in enumerate(sorted_dates_chrono, 9):
        ws_dash.row_dimensions[r_idx].height = 20
        rep = reports_obj.get(d_str, {})
        summ = rep.get('member', {}).get('summery', {})
        lot_summ = rep.get('lottery_sales_summary', {})

        b_tot = (summ.get('betAmount1DKhr', 0) or 0) + (summ.get('betAmount2DKhr', 0) or 0) + (summ.get('betAmount3DKhr', 0) or 0) + (summ.get('betAmount4DKhr', 0) or 0)
        net_tot = summ.get('winLoseAmountKhr', 0) or 0
        get_l_s = lambda code: (lot_summ.get(code) or {}).get('sales_khr', 0)

        row_vals = [d_str, get_l_s('MHSB'), get_l_s('MC'), get_l_s('MT'), get_l_s('KH'), get_l_s('TC'), get_l_s('SC'), get_l_s('TH'), get_l_s('KP'), b_tot, net_tot]
        for c_i, val in enumerate(row_vals, 1):
            cell = ws_dash.cell(row=r_idx, column=c_i, value=val)
            cell.font = cell_font
            cell.border = grid_border
            if c_i == 1:
                cell.alignment = align_center
                cell.font = blue_font
            else:
                cell.alignment = align_right
                if isinstance(val, (int, float)):
                    cell.number_format = '#,##0'
                    if val < 0:
                        cell.font = red_bold_font

    for col_i in range(1, 12):
        ws_dash.column_dimensions[get_column_letter(col_i)].width = 16
    ws_dash.column_dimensions['A'].width = 14
    ws_dash.column_dimensions['J'].width = 18
    ws_dash.column_dimensions['K'].width = 20

    # 4 Main Level Sheets
    levels_config = [
        {"title": "Member (សមាជិក)", "key": "member"},
        {"title": "Master (មេ)", "key": "master"},
        {"title": "Senior (សេនៀ)", "key": "senior"},
        {"title": "Super Senior (ស៊ុបភើ)", "key": "super_senior"}
    ]
    for lev in levels_config:
        ws = wb.create_sheet(title=lev["title"])
        ws.views.sheetView[0].showGridLines = True
        ws.row_dimensions[1].height = 28
        for col_idx, h_text in enumerate(HEADERS_32, 1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.fill = green_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = grid_border

        for d_str in sorted_dates_chrono:
            rep = reports_obj.get(d_str, {})
            settlements = rep.get(lev["key"], {}).get('settlements', [])
            append_rows_to_sheet(ws, d_str, "ALL", settlements, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right)

        for col_i in range(1, 33):
            ws.column_dimensions[get_column_letter(col_i)].width = 16
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 14
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 22

    # 8 Lottery Sheets
    for cfg in LOTTERY_CONFIGS:
        sheet_name = cfg["sheet"]
        l_code = cfg["code"]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True
        ws.row_dimensions[1].height = 28
        for col_idx, h_text in enumerate(HEADERS_32, 1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.fill = green_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = grid_border

        for d_str in sorted_dates_chrono:
            if d_str == target_date_display:
                settlements = lottery_settlements_map.get(l_code, [])
            else:
                # If cached from previous master
                settlements = []
            append_rows_to_sheet(ws, d_str, l_code, settlements, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right)

        for col_i in range(1, 33):
            ws.column_dimensions[get_column_letter(col_i)].width = 16
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 14
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 22

    # Sheet: លេខសាង
    ws_sang = wb.create_sheet(title="លេខសាង")
    ws_sang.views.sheetView[0].showGridLines = True
    ws_sang.row_dimensions[1].height = 28
    for col_idx, h_text in enumerate(HEADERS_32, 1):
        cell = ws_sang.cell(row=1, column=col_idx, value=h_text)
        cell.fill = green_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = grid_border

    for d_str in sorted_dates_chrono:
        rep = reports_obj.get(d_str, {})
        sang_sett = rep.get('super_senior_sang', {}).get('settlements', [])
        append_rows_to_sheet(ws_sang, d_str, "ALL", sang_sett, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right)

    for col_i in range(1, 33):
        ws_sang.column_dimensions[get_column_letter(col_i)].width = 16
    ws_sang.column_dimensions['A'].width = 14
    ws_sang.column_dimensions['B'].width = 14
    ws_sang.column_dimensions['C'].width = 18
    ws_sang.column_dimensions['D'].width = 22

    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    wb.save(master_excel_path)
    print(f"\n🎉 Successfully updated Master Excel: {master_excel_path}")

    # 5. Generate Standalone Daily Excel (total_summary_report_DD_MM_YYYY.xlsx)
    d_filename_str = target_date_display.replace('/', '_')
    daily_excel_filename = f"total_summary_report_{d_filename_str}.xlsx"
    daily_excel_path = os.path.join(os.getcwd(), daily_excel_filename)
    
    wb_daily = openpyxl.Workbook()
    def_sheet = wb_daily.active

    # Daily Dashboard
    ws_d_dash = wb_daily.create_sheet(title="Dashboard")
    ws_d_dash.views.sheetView[0].showGridLines = True
    ws_d_dash.merge_cells("A1:K2")
    ws_d_dash["A1"] = f"P99L LOTTERY SYSTEM — DAILY REPORT SUMMARY ({target_date_display})"
    ws_d_dash["A1"].fill = dark_green_fill
    ws_d_dash["A1"].font = title_font
    ws_d_dash["A1"].alignment = align_center

    m_summ = member_data.get('summery', {})
    d_sales_khr = (m_summ.get('betAmount1DKhr',0) or 0) + (m_summ.get('betAmount2DKhr',0) or 0) + (m_summ.get('betAmount3DKhr',0) or 0) + (m_summ.get('betAmount4DKhr',0) or 0)
    d_net_khr = m_summ.get('winLoseAmountKhr',0) or 0

    ws_d_dash.merge_cells("A4:E4")
    ws_d_dash["A4"] = "💰 លុយលក់សរុបប្រចាំថ្ងៃ (Daily Sales)"
    ws_d_dash["A4"].fill = green_fill
    ws_d_dash["A4"].font = header_font
    ws_d_dash["A4"].alignment = align_center
    ws_d_dash.merge_cells("A5:E5")
    ws_d_dash["A5"] = f"{round(d_sales_khr):,} ៛"
    ws_d_dash["A5"].font = Font(name="Kantumruy Pro", size=14, bold=True, color="2E7D32")
    ws_d_dash["A5"].alignment = align_center

    ws_d_dash.merge_cells("G4:K4")
    ws_d_dash["G4"] = "⚖️ ស៊ីខាតសរុបប្រចាំថ្ងៃ (Daily Net Win/Loss)"
    ws_d_dash["G4"].fill = green_fill
    ws_d_dash["G4"].font = header_font
    ws_d_dash["G4"].alignment = align_center
    ws_d_dash.merge_cells("G5:K5")
    ws_d_dash["G5"] = f"{round(d_net_khr):,} ៛"
    ws_d_dash["G5"].font = Font(name="Kantumruy Pro", size=14, bold=True, color="2E7D32" if d_net_khr >= 0 else "D93025")
    ws_d_dash["G5"].alignment = align_center

    ws_d_dash.merge_cells("A7:K7")
    ws_d_dash["A7"] = f"📊 តារាងសរុបស៊ីខាត និង លុយលក់តាមប្រភេទឆ្នោតនីមួយៗ ({target_date_display})"
    ws_d_dash["A7"].fill = dark_green_fill
    ws_d_dash["A7"].font = header_font
    ws_d_dash["A7"].alignment = align_left

    ws_d_dash.row_dimensions[8].height = 24
    for c_i, h_t in enumerate(dash_headers, 1):
        cell = ws_d_dash.cell(row=8, column=c_i, value=h_t)
        cell.fill = green_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = grid_border

    get_l_s = lambda code: (lottery_sales_summary.get(code) or {}).get('sales_khr', 0)
    d_row_vals = [target_date_display, get_l_s('MHSB'), get_l_s('MC'), get_l_s('MT'), get_l_s('KH'), get_l_s('TC'), get_l_s('SC'), get_l_s('TH'), get_l_s('KP'), d_sales_khr, d_net_khr]
    ws_d_dash.row_dimensions[9].height = 20
    for c_i, val in enumerate(d_row_vals, 1):
        cell = ws_d_dash.cell(row=9, column=c_i, value=val)
        cell.font = cell_font
        cell.border = grid_border
        if c_i == 1:
            cell.alignment = align_center
            cell.font = blue_font
        else:
            cell.alignment = align_right
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0'
                if val < 0:
                    cell.font = red_bold_font

    for col_i in range(1, 12):
        ws_d_dash.column_dimensions[get_column_letter(col_i)].width = 16
    ws_d_dash.column_dimensions['A'].width = 14
    ws_d_dash.column_dimensions['J'].width = 18
    ws_d_dash.column_dimensions['K'].width = 20

    # 4 Main Level Sheets for Daily Excel
    for lev in levels_config:
        ws_l = wb_daily.create_sheet(title=lev["title"])
        ws_l.views.sheetView[0].showGridLines = True
        ws_l.row_dimensions[1].height = 28
        for col_idx, h_text in enumerate(HEADERS_32, 1):
            cell = ws_l.cell(row=1, column=col_idx, value=h_text)
            cell.fill = green_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = grid_border

        target_rep = reports_obj.get(target_date_display, {})
        settlements = target_rep.get(lev["key"], {}).get('settlements', [])
        append_rows_to_sheet(ws_l, target_date_display, "ALL", settlements, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right)

        for col_i in range(1, 33):
            ws_l.column_dimensions[get_column_letter(col_i)].width = 16
        ws_l.column_dimensions['A'].width = 14
        ws_l.column_dimensions['B'].width = 14
        ws_l.column_dimensions['C'].width = 18
        ws_l.column_dimensions['D'].width = 22

    # 8 Lottery Sheets for Daily Excel
    for cfg in LOTTERY_CONFIGS:
        sheet_name = cfg["sheet"]
        l_code = cfg["code"]
        ws_cfg = wb_daily.create_sheet(title=sheet_name)
        ws_cfg.views.sheetView[0].showGridLines = True
        ws_cfg.row_dimensions[1].height = 28
        for col_idx, h_text in enumerate(HEADERS_32, 1):
            cell = ws_cfg.cell(row=1, column=col_idx, value=h_text)
            cell.fill = green_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = grid_border

        settlements = lottery_settlements_map.get(l_code, [])
        append_rows_to_sheet(ws_cfg, target_date_display, l_code, settlements, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right)

        for col_i in range(1, 33):
            ws_cfg.column_dimensions[get_column_letter(col_i)].width = 16
        ws_cfg.column_dimensions['A'].width = 14
        ws_cfg.column_dimensions['B'].width = 14
        ws_cfg.column_dimensions['C'].width = 18
        ws_cfg.column_dimensions['D'].width = 22

    # Sang Sheet for Daily Excel
    ws_d_sang = wb_daily.create_sheet(title="លេខសាង")
    ws_d_sang.views.sheetView[0].showGridLines = True
    ws_d_sang.row_dimensions[1].height = 28
    for col_idx, h_text in enumerate(HEADERS_32, 1):
        cell = ws_d_sang.cell(row=1, column=col_idx, value=h_text)
        cell.fill = green_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = grid_border

    sang_sett = super_senior_sang_data.get('settlements', [])
    append_rows_to_sheet(ws_d_sang, target_date_display, "ALL", sang_sett, green_fill, header_font, cell_font, blue_font, red_bold_font, grid_border, align_center, align_left, align_right)

    for col_i in range(1, 33):
        ws_d_sang.column_dimensions[get_column_letter(col_i)].width = 16
    ws_d_sang.column_dimensions['A'].width = 14
    ws_d_sang.column_dimensions['B'].width = 14
    ws_d_sang.column_dimensions['C'].width = 18
    ws_d_sang.column_dimensions['D'].width = 22

    if def_sheet in wb_daily.worksheets:
        wb_daily.remove(def_sheet)

    wb_daily.save(daily_excel_path)
    print(f"🎉 Successfully created Daily Standalone Excel: {daily_excel_path}")

    try:
        from export_monthly_matrix import build_monthly_excel
        build_monthly_excel()
    except Exception as e:
        print(f"Note: Monthly excel update skipped: {e}")

    # 6. Auto-sync to Google Drive folder
    try:
        import subprocess
        subprocess.run(["python3", "prepare_google_drive_folder.py"], check=True)
        print("☁️ Auto-synced all new reports to Google_Drive_P99L folder successfully!")
    except Exception as e:
        print(f"⚠️ Google Drive sync skipped: {e}")

    # 7. Auto-push to GitHub Pages for 24/7 Permanent Web Portal
    try:
        import subprocess
        subprocess.run(["git", "add", "index.html", "styles.css", "app.js", "total_summary_report.json"], check=False)
        subprocess.run(["git", "commit", "-m", f"Auto-update reports for date {target_date_display}"], check=False)
        subprocess.run(["git", "push", "origin", "main"], check=False)
        print("🚀 Auto-pushed latest data to GitHub Pages (Live 24/7 Web Portal) successfully!")
    except Exception as e:
        print(f"⚠️ GitHub Pages auto-push skipped: {e}")

if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else None
    update_daily_report(target)
