#!/usr/bin/env python3
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MONTH_MAP = {
    '01': ('january', 'មករា', 'JANUARY'),
    '02': ('february', 'កុម្ភៈ', 'FEBRUARY'),
    '03': ('march', 'មីនា', 'MARCH'),
    '04': ('april', 'មេសា', 'APRIL'),
    '05': ('may', 'ឧសភា', 'MAY'),
    '06': ('june', 'មិថុនា', 'JUNE'),
    '07': ('july', 'កក្កដា', 'JULY'),
    '08': ('august', 'សីហា', 'AUGUST'),
    '09': ('september', 'កញ្ញា', 'SEPTEMBER'),
    '10': ('october', 'តុលា', 'OCTOBER'),
    '11': ('november', 'វិច្ឆិកា', 'NOVEMBER'),
    '12': ('december', 'ធ្នូ', 'DECEMBER')
}

def build_monthly_excel(json_path='total_summary_report.json', output_file=None):
    if not os.path.exists(json_path):
        print(f'File {json_path} not found.')
        return None

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    reports = data.get('reports', {})
    dates = data.get('dates_available', list(reports.keys()))
    
    # Determine Month and Year dynamically from dates
    m_en, m_kh, m_upper = 'september', 'កញ្ញា', 'SEPTEMBER'
    y_str = '2026'
    if dates:
        first_d = dates[0] # e.g. '01/09/2026'
        parts = first_d.split('/')
        if len(parts) == 3:
            m_num = parts[1]
            y_str = parts[2]
            m_en, m_kh, m_upper = MONTH_MAP.get(m_num, ('september', 'កញ្ញា', 'SEPTEMBER'))

    if not output_file:
        output_file = f'monthly_net_summary_{m_en}_{y_str}.xlsx'
    
    LOTTERIES = [
        ('MHSB', 'លាភមហាសម្បត្តិ (LMHSB)'),
        ('MT',   'មហាទេព (MT)'),
        ('TC',   'ទិញឈ្នះ (TC)'),
        ('MC',   'មហាឈ្នះ (MC)'),
        ('SC',   'សប្បាយឈ្នះ (SC)'),
        ('KH',   'ឆ្នោតខ្មែរ (KH)'),
        ('TH',   'ឆ្នោតថៃ (TH)'),
        ('KP',   'ឆ្នោតកំពត (KP)')
    ]

    ss_map = {}
    for d_str, rep in reports.items():
        sett = rep.get('super_senior', {}).get('settlements', [])
        for item in sett:
            u = item.get('username') or item.get('userCode')
            if u and u not in ss_map:
                ss_map[u] = item.get('nickname', '')

    sorted_ss = sorted(list(ss_map.keys()))

    monthly_rows = []
    col_totals = {code: {'khr': 0.0, 'usd': 0.0} for code, _ in LOTTERIES}
    grand_total_khr = 0.0
    grand_total_usd = 0.0

    for idx, u in enumerate(sorted_ss, start=1):
        nick = ss_map[u]
        row = {
            'no': idx,
            'username': u,
            'nickname': nick,
            'lots': {},
            'tot_khr': 0.0,
            'tot_usd': 0.0
        }
        for code, name in LOTTERIES:
            khr_sum = 0.0
            usd_sum = 0.0
            for d_str in dates:
                ss_lot = reports.get(d_str, {}).get('super_senior_lotteries', {}).get(code, {})
                u_data = ss_lot.get(u, {})
                khr_sum += float(u_data.get('net_khr', 0) or 0)
                usd_sum += float(u_data.get('net_usd', 0) or 0)
            row['lots'][code] = {'khr': khr_sum, 'usd': usd_sum}
            row['tot_khr'] += khr_sum
            row['tot_usd'] += usd_sum
            col_totals[code]['khr'] += khr_sum
            col_totals[code]['usd'] += usd_sum

        grand_total_khr += row['tot_khr']
        grand_total_usd += row['tot_usd']
        monthly_rows.append(row)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Monthly Summary 100%'

    dark_green_fill = PatternFill(start_color='1B5E20', end_color='1B5E20', fill_type='solid')
    med_green_fill = PatternFill(start_color='2E7D32', end_color='2E7D32', fill_type='solid')
    light_total_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
    grand_th_fill = PatternFill(start_color='14532D', end_color='14532D', fill_type='solid')

    font_title = Font(name='Kantumruy Pro', size=16, bold=True, color='1B5E20')
    font_th1 = Font(name='Kantumruy Pro', size=11, bold=True, color='FFFFFF')
    font_th2 = Font(name='Outfit', size=10, bold=True, color='FFFFFF')
    font_grand_th = Font(name='Kantumruy Pro', size=11, bold=True, color='FEF08A')
    font_data = Font(name='Outfit', size=10)
    font_data_bold = Font(name='Outfit', size=10, bold=True)
    font_neg = Font(name='Outfit', size=10, bold=True, color='DC2626')
    font_zero = Font(name='Outfit', size=10, color='94A3B8')
    font_tfoot = Font(name='Outfit', size=11, bold=True, color='FFFFFF')

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    group_right_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='medium', color='1B5E20'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    tfoot_border = Border(
        left=Side(style='thin', color='2E7D32'),
        right=Side(style='thin', color='2E7D32'),
        top=Side(style='double', color='FFFFFF'),
        bottom=Side(style='double', color='FFFFFF')
    )

    ws.merge_cells('A1:V1')
    cell_t = ws['A1']
    cell_t.value = f'តារាងស៊ីខាតសរុបប្រចាំខែ 100% — {m_upper}-{y_str}'
    cell_t.font = font_title
    cell_t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:A3')
    ws['A2'] = 'NO'
    ws['A2'].fill = dark_green_fill
    ws['A2'].font = font_th1
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.merge_cells('B2:B3')
    ws['B2'] = 'Super Senior'
    ws['B2'].fill = dark_green_fill
    ws['B2'].font = font_th1
    ws['B2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.merge_cells('C2:C3')
    ws['C2'] = 'Nick Name'
    ws['C2'].fill = dark_green_fill
    ws['C2'].font = font_th1
    ws['C2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    curr_col = 4
    for code, name in LOTTERIES:
        c_khr = curr_col
        c_usd = curr_col + 1
        
        ws.merge_cells(start_row=2, start_column=c_khr, end_row=2, end_column=c_usd)
        th1 = ws.cell(row=2, column=c_khr, value=name)
        th1.fill = dark_green_fill
        th1.font = font_th1
        th1.alignment = Alignment(horizontal='center', vertical='center')
        
        th2_k = ws.cell(row=3, column=c_khr, value='(KHR)')
        th2_k.fill = med_green_fill
        th2_k.font = font_th2
        th2_k.alignment = Alignment(horizontal='center', vertical='center')
        
        th2_u = ws.cell(row=3, column=c_usd, value='(USD)')
        th2_u.fill = med_green_fill
        th2_u.font = font_th2
        th2_u.alignment = Alignment(horizontal='center', vertical='center')
        
        curr_col += 2

    c_tot_k = curr_col
    c_tot_u = curr_col + 1
    ws.merge_cells(start_row=2, start_column=c_tot_k, end_row=2, end_column=c_tot_u)
    th_gt = ws.cell(row=2, column=c_tot_k, value='សរុប (Grand Total)')
    th_gt.fill = grand_th_fill
    th_gt.font = font_grand_th
    th_gt.alignment = Alignment(horizontal='center', vertical='center')

    th_gt_k = ws.cell(row=3, column=c_tot_k, value='(KHR)')
    th_gt_k.fill = med_green_fill
    th_gt_k.font = font_th2
    th_gt_k.alignment = Alignment(horizontal='center', vertical='center')

    th_gt_u = ws.cell(row=3, column=c_tot_u, value='(USD)')
    th_gt_u.fill = med_green_fill
    th_gt_u.font = font_th2
    th_gt_u.alignment = Alignment(horizontal='center', vertical='center')

    curr_col += 2
    c_ss_r = curr_col
    ws.merge_cells(start_row=2, start_column=c_ss_r, end_row=3, end_column=c_ss_r)
    th_ssr = ws.cell(row=2, column=c_ss_r, value='Super Senior')
    th_ssr.fill = dark_green_fill
    th_ssr.font = font_th1
    th_ssr.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 20

    row_idx = 4
    for r in monthly_rows:
        ws.row_dimensions[row_idx].height = 19
        
        c = ws.cell(row=row_idx, column=1, value=r['no'])
        c.font = font_data_bold
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        
        c = ws.cell(row=row_idx, column=2, value=r['username'])
        c.font = font_data_bold
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        
        c = ws.cell(row=row_idx, column=3, value=r['nickname'])
        c.font = Font(name='Kantumruy Pro', size=10)
        c.alignment = Alignment(horizontal='left', vertical='center')
        c.border = group_right_border

        col_pos = 4
        for code, _ in LOTTERIES:
            k_val = r['lots'][code]['khr']
            u_val = r['lots'][code]['usd']

            ck = ws.cell(row=row_idx, column=col_pos, value=k_val)
            ck.number_format = '#,##0"៛";[Red]-#,##0"៛";"-"'
            ck.alignment = Alignment(horizontal='right', vertical='center')
            ck.border = thin_border
            ck.font = font_neg if k_val < 0 else (font_zero if k_val == 0 else font_data)

            cu = ws.cell(row=row_idx, column=col_pos + 1, value=u_val)
            cu.number_format = '0,##0.00;[Red]-0,##0.00;"$ -"'
            cu.alignment = Alignment(horizontal='right', vertical='center')
            cu.border = group_right_border
            cu.font = font_neg if u_val < 0 else (font_zero if u_val == 0 else font_data)

            col_pos += 2

        ck_gt = ws.cell(row=row_idx, column=col_pos, value=r['tot_khr'])
        ck_gt.number_format = '#,##0"៛";[Red]-#,##0"៛";"-"'
        ck_gt.alignment = Alignment(horizontal='right', vertical='center')
        ck_gt.fill = light_total_fill
        ck_gt.border = thin_border
        ck_gt.font = font_neg if r['tot_khr'] < 0 else font_data_bold

        cu_gt = ws.cell(row=row_idx, column=col_pos + 1, value=r['tot_usd'])
        cu_gt.number_format = '0,##0.00;[Red]-0,##0.00;"$ -"'
        cu_gt.alignment = Alignment(horizontal='right', vertical='center')
        cu_gt.fill = light_total_fill
        cu_gt.border = group_right_border
        cu_gt.font = font_neg if r['tot_usd'] < 0 else font_data_bold

        col_pos += 2
        css_r = ws.cell(row=row_idx, column=col_pos, value=r['username'])
        css_r.font = font_data_bold
        css_r.alignment = Alignment(horizontal='center', vertical='center')
        css_r.border = thin_border

        row_idx += 1

    ws.row_dimensions[row_idx].height = 24
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
    c_tot_lbl = ws.cell(row=row_idx, column=1, value='TOTAL')
    c_tot_lbl.fill = dark_green_fill
    c_tot_lbl.font = Font(name='Outfit', size=12, bold=True, color='FFFFFF')
    c_tot_lbl.alignment = Alignment(horizontal='center', vertical='center')
    c_tot_lbl.border = tfoot_border

    col_pos = 4
    for code, _ in LOTTERIES:
        k_tot = col_totals[code]['khr']
        u_tot = col_totals[code]['usd']

        ck = ws.cell(row=row_idx, column=col_pos, value=k_tot)
        ck.number_format = '#,##0"៛";[Red]-#,##0"៛";"-"'
        ck.fill = dark_green_fill
        ck.font = font_tfoot
        ck.alignment = Alignment(horizontal='right', vertical='center')
        ck.border = tfoot_border

        cu = ws.cell(row=row_idx, column=col_pos + 1, value=u_tot)
        cu.number_format = '0,##0.00;[Red]-0,##0.00;"$ -"'
        cu.fill = dark_green_fill
        cu.font = font_tfoot
        cu.alignment = Alignment(horizontal='right', vertical='center')
        cu.border = tfoot_border

        col_pos += 2

    ck_gt = ws.cell(row=row_idx, column=col_pos, value=grand_total_khr)
    ck_gt.number_format = '#,##0"៛";[Red]-#,##0"៛";"-"'
    ck_gt.fill = grand_th_fill
    ck_gt.font = Font(name='Outfit', size=11, bold=True, color='FEF08A')
    ck_gt.alignment = Alignment(horizontal='right', vertical='center')
    ck_gt.border = tfoot_border

    cu_gt = ws.cell(row=row_idx, column=col_pos + 1, value=grand_total_usd)
    cu_gt.number_format = '0,##0.00;[Red]-0,##0.00;"$ -"'
    cu_gt.fill = grand_th_fill
    cu_gt.font = Font(name='Outfit', size=11, bold=True, color='FEF08A')
    cu_gt.alignment = Alignment(horizontal='right', vertical='center')
    cu_gt.border = tfoot_border

    col_pos += 2
    c_tot_r = ws.cell(row=row_idx, column=col_pos, value='TOTAL')
    c_tot_r.fill = dark_green_fill
    c_tot_r.font = Font(name='Outfit', size=11, bold=True, color='FFFFFF')
    c_tot_r.alignment = Alignment(horizontal='center', vertical='center')
    c_tot_r.border = tfoot_border

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 18
    for col_idx in range(4, 22):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 15
    ws.column_dimensions['V'].width = 14

    wb.save(output_file)
    print(f'✅ Successfully generated Monthly Summary Excel: {output_file}')
    return output_file

if __name__ == '__main__':
    build_monthly_excel()
